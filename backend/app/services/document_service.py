from dataclasses import dataclass
from io import BytesIO
import math
from pathlib import Path
import re
import zipfile

from app.config import get_settings


class DocumentError(Exception):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


@dataclass(frozen=True)
class PageText:
    page: int
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    pages: list[PageText]
    page_count: int | None
    extraction_method: str

    @property
    def text(self) -> str:
        return "\n\n".join(
            f"[Page {page.page}]\n{page.text.strip()}" for page in self.pages if page.text.strip()
        )


def validate_file_content(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".docx", ".txt"}:
        raise DocumentError("Unsupported file type. Upload a PDF, DOCX, or TXT file.", 415)
    if not content:
        raise DocumentError("The uploaded file is empty.", 400)
    if suffix == ".pdf" and not content.startswith(b"%PDF-"):
        raise DocumentError("The file content does not match a PDF document.", 415)
    if suffix == ".docx":
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise DocumentError("The file content is not a valid DOCX document.", 415)
        except zipfile.BadZipFile as exc:
            raise DocumentError("The file content is not a valid DOCX document.", 415) from exc
    if suffix == ".txt":
        if b"\x00" in content:
            raise DocumentError("The TXT file contains binary data.", 415)
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("cp1252")
            except UnicodeDecodeError as exc:
                raise DocumentError("The TXT file could not be decoded as text.", 422) from exc
        if not text.strip():
            raise DocumentError("The uploaded file contains no readable text.", 422)
        control_count = sum(ord(char) < 32 and char not in "\n\r\t\f" for char in text)
        if control_count / max(len(text), 1) > 0.02:
            raise DocumentError("The TXT file contains unreadable binary data.", 415)
    return suffix


def extract_document(filename: str, content: bytes) -> ExtractedDocument:
    suffix = validate_file_content(filename, content)
    if suffix == ".pdf":
        return _extract_pdf(content)
    if suffix == ".docx":
        return _extract_docx(content)
    text = content.decode("utf-8-sig", errors="replace")
    if "\ufffd" in text:
        text = content.decode("cp1252")
    if not text.strip():
        raise DocumentError("The uploaded file contains no readable text.", 422)
    return ExtractedDocument([PageText(1, text.strip())], None, "text")


def _extract_pdf(content: bytes) -> ExtractedDocument:
    try:
        import fitz
    except ImportError as exc:
        raise DocumentError("PDF support requires PyMuPDF. Install the backend requirements.", 503) from exc
    try:
        document = fitz.open(stream=content, filetype="pdf")
    except (RuntimeError, ValueError) as exc:
        raise DocumentError("The PDF could not be opened or is not readable.", 422) from exc
    with document:
        settings = get_settings()
        if document.page_count == 0:
            raise DocumentError("The PDF contains no pages.", 422)
        if document.page_count > settings.max_pdf_pages:
            raise DocumentError(f"The PDF exceeds the {settings.max_pdf_pages}-page processing limit.", 413)
        extracted: list[PageText] = []
        pages_needing_ocr = []
        try:
            for page_number, page in enumerate(document, start=1):
                page_text = page.get_text("text").strip()
                extracted.append(PageText(page_number, page_text))
                if len(page_text) < settings.min_page_text_chars:
                    pages_needing_ocr.append(page_number - 1)
        except (RuntimeError, ValueError) as exc:
            raise DocumentError("The PDF could not be read or is password-protected.", 422) from exc

        text_chars = sum(len(item.text) for item in extracted)
        if text_chars < settings.min_pdf_text_chars and not pages_needing_ocr:
            pages_needing_ocr = list(range(document.page_count))
        if not pages_needing_ocr:
            if text_chars == 0:
                raise DocumentError("The PDF contains no extractable text.", 422)
            return ExtractedDocument(extracted, document.page_count, "pdf_text")

        try:
            import pytesseract
            from PIL import Image
        except ImportError as exc:
            raise DocumentError(
                "OCR is required for this scanned PDF. Install pytesseract, Pillow, and the Tesseract executable.",
                503,
            ) from exc
        try:
            for page_index in pages_needing_ocr:
                rect = document[page_index].rect
                width_points = float(rect.width)
                height_points = float(rect.height)
                if width_points <= 0 or height_points <= 0:
                    raise DocumentError("A PDF page has invalid dimensions.", 422)
                dpi_for_dimension = (
                    settings.max_ocr_image_dimension * 72 / max(width_points, height_points)
                )
                dpi_for_pixels = math.sqrt(
                    settings.max_ocr_image_pixels * 72**2 / (width_points * height_points)
                )
                effective_dpi = min(settings.ocr_dpi, dpi_for_dimension, dpi_for_pixels)
                if effective_dpi < 72:
                    raise DocumentError(
                        "A PDF page exceeds the safe OCR image-resolution limit.", 413
                    )
                pixmap = document[page_index].get_pixmap(
                    dpi=int(effective_dpi), alpha=False, colorspace=fitz.csRGB
                )
                with Image.open(BytesIO(pixmap.tobytes("png"))) as image:
                    ocr_text = pytesseract.image_to_string(image).strip()
                if ocr_text:
                    extracted[page_index] = PageText(page_index + 1, ocr_text)
        except pytesseract.pytesseract.TesseractNotFoundError as exc:
            raise DocumentError("Tesseract OCR is not installed or is not available on PATH.", 503) from exc
        except (RuntimeError, OSError) as exc:
            raise DocumentError("OCR could not read one or more scanned PDF pages.", 422) from exc
        if not any(item.text.strip() for item in extracted):
            raise DocumentError("No readable text was found in the PDF, including through OCR.", 422)
        method = "pdf_ocr" if text_chars < settings.min_pdf_text_chars else "pdf_text_and_ocr"
        return ExtractedDocument(extracted, document.page_count, method)


def _extract_docx(content: bytes) -> ExtractedDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentError("DOCX support requires python-docx. Install the backend requirements.", 503) from exc
    try:
        document = Document(BytesIO(content))
    except (ValueError, zipfile.BadZipFile) as exc:
        raise DocumentError("The DOCX could not be opened or is not readable.", 422) from exc
    paragraphs = [re.sub(r"\s+", " ", paragraph.text).strip() for paragraph in document.paragraphs]
    text = "\n".join(paragraph for paragraph in paragraphs if paragraph)
    if not text:
        raise DocumentError("The DOCX contains no readable text.", 422)
    return ExtractedDocument([PageText(1, text)], None, "docx")
