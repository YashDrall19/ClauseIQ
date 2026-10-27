from io import BytesIO
from types import SimpleNamespace
import sys

import pytest
from docx import Document

from app.services.document_service import DocumentError, extract_document, validate_file_content


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Payment is due within thirty days.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf(text: str) -> bytes:
    fitz = pytest.importorskip("fitz")
    document = fitz.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def test_extract_txt():
    result = extract_document("terms.txt", b"First party pays the second party.")
    assert result.extraction_method == "text"
    assert result.pages[0].page == 1
    assert "First party" in result.text


def test_extract_docx():
    result = extract_document("terms.docx", make_docx())
    assert result.extraction_method == "docx"
    assert "thirty days" in result.text


def test_extract_pdf_preserves_page_numbers():
    result = extract_document(
        "terms.pdf",
        make_pdf("This is a readable contract page with enough extractable text. " * 3),
    )
    assert result.extraction_method == "pdf_text"
    assert result.page_count == 1
    assert result.pages[0].page == 1


def test_pdf_text_extraction_preserves_page_numbers_with_pymupdf(monkeypatch):
    page = SimpleNamespace(
        rect=SimpleNamespace(width=612, height=792),
        get_text=lambda _kind: "Readable contract terms and conditions. " * 4,
    )

    class FakeDocument:
        page_count = 1

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

        def __iter__(self):
            return iter([page])

    fake_fitz = SimpleNamespace(open=lambda **_kwargs: FakeDocument(), csRGB="RGB")
    monkeypatch.setitem(sys.modules, "fitz", fake_fitz)
    result = extract_document("terms.pdf", b"%PDF-fake")
    assert result.extraction_method == "pdf_text"
    assert result.page_count == 1
    assert result.pages[0].page == 1


def test_pdf_ocr_failure_is_reported(monkeypatch, test_settings):
    import pytesseract

    test_settings.min_page_text_chars = 100
    image_buffer = BytesIO()
    from PIL import Image

    Image.new("RGB", (1, 1)).save(image_buffer, format="PNG")

    class FakePixmap:
        def tobytes(self, _format):
            return image_buffer.getvalue()

    class FakePage:
        rect = SimpleNamespace(width=612, height=792)

        def get_text(self, _kind):
            return ""

        def get_pixmap(self, **_kwargs):
            return FakePixmap()

    class FakeDocument:
        page_count = 1

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

        def __iter__(self):
            return iter([FakePage()])

        def __getitem__(self, _index):
            return FakePage()

    monkeypatch.setitem(
        sys.modules,
        "fitz",
        SimpleNamespace(open=lambda **_kwargs: FakeDocument(), csRGB="RGB"),
    )

    def no_tesseract(_image):
        raise pytesseract.pytesseract.TesseractNotFoundError()

    monkeypatch.setattr(pytesseract, "image_to_string", no_tesseract)
    with pytest.raises(DocumentError, match="Tesseract OCR is not installed"):
        extract_document("scan.pdf", b"%PDF-fake")


def test_reject_unsupported_and_empty_files():
    with pytest.raises(DocumentError) as unsupported:
        validate_file_content("contract.exe", b"data")
    assert unsupported.value.status_code == 415
    with pytest.raises(DocumentError) as empty:
        validate_file_content("contract.txt", b"")
    assert empty.value.status_code == 400


def test_reject_content_mismatched_to_extension():
    with pytest.raises(DocumentError, match="does not match"):
        validate_file_content("contract.pdf", b"not a pdf")

